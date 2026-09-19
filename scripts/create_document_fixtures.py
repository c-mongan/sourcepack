"""Author synthetic document fixtures. Optional authoring libs, never user evidence."""
import sys
from pathlib import Path
from docx import Document
from pptx import Presentation
from openpyxl import Workbook
from reportlab.pdfgen.canvas import Canvas
out=Path(sys.argv[1]);out.mkdir(parents=True,exist_ok=True)
doc=Document();doc.add_heading('Retry policy',0);doc.add_paragraph('The retry limit is seven. A human reviews failures.');doc.save(out/'policy.docx')
pres=Presentation()
for title,body in [('Setup','Set review mode to manual.'),('Handoff','The librarian receives approved work only.')]:
 slide=pres.slides.add_slide(pres.slide_layouts[1]);slide.shapes.title.text=title;slide.placeholders[1].text=body
pres.save(out/'handoff.pptx')
book=Workbook();ws=book.active;ws.title='Budgets';ws.append(['Tool','Required','Budget']);ws.append(['Reader','Yes',7]);ws.append(['OCR','No',0]);book.save(out/'budgets.xlsx')
pdf=Canvas(str(out/'policy.pdf'));pdf.drawString(72,740,'The retry limit is seven.');pdf.drawString(72,710,'Human review is required before publication.');pdf.showPage();pdf.save()
# Layout/OCR fixture intentionally requires visual reading; no text layer.
from PIL import Image,ImageDraw
img=Image.new('RGB',(1200,800),'white');draw=ImageDraw.Draw(img);draw.text((80,80),'VISUAL SETTING: REVIEW = MANUAL',fill='black',font_size=44);img.save(out/'scan.pdf','PDF')
print(out)
